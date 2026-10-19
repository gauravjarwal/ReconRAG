from __future__ import annotations

import io
import logging
from collections.abc import Callable

logger = logging.getLogger(__name__)


class DocumentParseError(ValueError):
    """The file could not be turned into text. Message is safe to show the client."""


def parse_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise DocumentParseError("PDF is encrypted")
        pages = [page.extract_text() or "" for page in reader.pages]
    except PdfReadError as exc:
        raise DocumentParseError("PDF is corrupt or unreadable") from exc
    return "\n".join(pages)


def parse_docx(data: bytes) -> str:
    import docx

    try:
        doc = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises zipfile/KeyError/lxml errors on bad input
        raise DocumentParseError("DOCX is corrupt or unreadable") from exc
    parts = [para.text for para in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def parse_text(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


PARSERS: dict[str, Callable[[bytes], str]] = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".txt": parse_text,
    ".md": parse_text,
}

ALLOWED_EXTENSIONS = set(PARSERS.keys())


def file_extension(filename: str) -> str:
    return "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def parse_file(filename: str, data: bytes) -> str:
    """Extract raw text from a file given its name and bytes."""
    ext = file_extension(filename)
    parser = PARSERS.get(ext)
    if parser is None:
        raise DocumentParseError(f"Unsupported file type: {ext!r}")
    text = parser(data)
    logger.debug("Parsed %s file — %d chars", ext, len(text))
    return text
